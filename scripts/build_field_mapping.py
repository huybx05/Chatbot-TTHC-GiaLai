"""Sinh bảng A.2, field_mapping.csv và .xlsx từ dữ liệu crawl thật.

Nguồn chân lý: data/json/ — mỗi file có `domain_group` (suy từ tên sheet) và `linh_vuc`.
canonical_field đọc từ config/field_mapping.csv hiện có (giữ nguyên, không đổi).
"""

from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

REPO = Path(__file__).resolve().parents[1]
CSV = REPO / "config" / "field_mapping.csv"
XLSX = REPO / "config" / "field_mapping.xlsx"

COLS = ["domain_group", "source_field", "canonical_field", "so_tthc", "review_status"]

# canonical_field cho các lĩnh vực CHƯA có trong mapping cũ. Đây là quyết định nghiệp vụ,
# không suy ra được từ dữ liệu. 4 TTHC liên thông không lấy được qua API nên phải khai báo.
CANONICAL_BO_SUNG = {
    "Thu BHXH, BHYT, BHTN, BH tai nạn lao động, BNN": "AN_SINH_XA_HOI",
    "Hộ tịch": "HO_TICH_QUOC_TICH",
    "Văn hóa": "DI_SAN_NGHE_THUAT_VAN_HOA",
}

# Giá trị kỳ vọng — sai lệch thì script báo để không âm thầm ghi số sai.
TONG_TTHC_KY_VONG = 1955
# 169 cặp (domain_group, source_field) = 168 lĩnh vực nguồn + "An toàn thực phẩm" ở 2 nhóm
EXPECTED_PAIR_COUNT = 169
SO_LINH_VUC_KY_VONG = 168


def doc_du_lieu() -> tuple[Counter, dict[str, str]]:
    """Đếm (nhóm, lĩnh vực) và giữ ánh xạ canonical đã duyệt."""
    dem: Counter = Counter()
    for p in (REPO / "data/json").rglob("*.json"):
        d = json.loads(p.read_text(encoding="utf-8"))
        nhom = (d.get("domain_group") or "").strip()
        lv = (d.get("linh_vuc") or "").strip()
        if nhom and lv:
            dem[(nhom, lv)] += 1

    canonical: dict[str, str] = {}
    if CSV.exists():
        for r in csv.DictReader(CSV.open(encoding="utf-8-sig")):
            if r.get("source_field") and r.get("canonical_field"):
                canonical.setdefault(r["source_field"].strip(), r["canonical_field"].strip())

    for lv, c in CANONICAL_BO_SUNG.items():
        canonical.setdefault(lv, c)

    return dem, canonical


def main() -> int:
    dem, canonical = doc_du_lieu()

    thieu_canonical = sorted({lv for _nhom, lv in dem} - set(canonical))
    if thieu_canonical:
        print("❌ Lĩnh vực chưa có canonical_field:")
        for lv in thieu_canonical:
            print(f"   - {lv!r}")
        return 1

    records = []
    for (nhom, lv), so in sorted(dem.items()):
        records.append(
            {
                "domain_group": nhom,
                "source_field": lv,
                "canonical_field": canonical[lv],
                "so_tthc": so,
                "review_status": "",
            }
        )

    tong = sum(r["so_tthc"] for r in records)
    pair_count = len(records)
    so_lv = len({r["source_field"] for r in records})
    so_canonical = len({r["canonical_field"] for r in records})
    print(f"dòng              : {pair_count}")
    print(f"domain_group      : {len({r['domain_group'] for r in records})}")
    print(f"source_field      : {so_lv}")
    print(f"canonical_field   : {so_canonical}")
    print(f"tổng TTHC         : {tong}")

    if tong != TONG_TTHC_KY_VONG:
        print(f"  ⚠️  Tổng TTHC lệch (kỳ vọng {TONG_TTHC_KY_VONG})")
    if pair_count != EXPECTED_PAIR_COUNT:
        print(f"  ⚠️  Số cặp lệch (kỳ vọng {EXPECTED_PAIR_COUNT})")
    if so_lv != SO_LINH_VUC_KY_VONG:
        print(f"  ⚠️  Số lĩnh vực lệch (kỳ vọng {SO_LINH_VUC_KY_VONG})")

    # --- CSV ---
    with CSV.open("w", encoding="utf-8-sig", newline="") as h:
        w = csv.DictWriter(h, fieldnames=COLS)
        w.writeheader()
        w.writerows(records)
    print(f"\nĐã ghi {CSV.relative_to(REPO)}")

    # --- XLSX ---
    wb = Workbook()
    ws = wb.active
    ws.title = "field_mapping"
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5496")
    for i, ten in enumerate(COLS, 1):
        o = ws.cell(row=1, column=i, value=ten)
        o.font, o.fill = hf, hfill
        o.alignment = Alignment(horizontal="center", vertical="center")
    for r_i, rec in enumerate(records, 2):
        for c_i, ten in enumerate(COLS, 1):
            ws.cell(row=r_i, column=c_i, value=rec[ten])
    for c_i, wd in enumerate([32, 46, 32, 10, 14], 1):
        ws.column_dimensions[get_column_letter(c_i)].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(records) + 1}"
    wb.save(XLSX)
    print(f"Đã ghi {XLSX.relative_to(REPO)}")

    # --- Xuất bảng markdown A.2 ---
    dong = [
        "| # | Mã Thư mục (`Domain_Group`) | Lĩnh vực Nguồn (`source_field`) "
        "| Canonical Field | Số TTHC |",
        "|:---:|---|---|---|:---:|",
    ]
    for r in records:
        so_hieu = r["domain_group"].split(".")[0].lstrip("0")
        dong.append(
            f"| {so_hieu} | `{r['domain_group']}` | {r['source_field']} "
            f"| `{r['canonical_field']}` | {r['so_tthc']} |"
        )
    (REPO / ".tmp_a2_only.md").write_text("\n".join(dong) + "\n", encoding="utf-8")
    print(f"Đã ghi .tmp_a2_only.md ({len(records)} dòng)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
